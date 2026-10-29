"""PowerPoint tools (ARCHITECTURE.md §8.2): read, and update by ops."""

from __future__ import annotations

import asyncio
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pptx

from deskai.tools.base import (
    RiskLevel,
    Tool,
    ToolContext,
    ToolResult,
    as_int,
    backup_existing,
    truncate,
)


class ReadPptxTool(Tool):
    name = "read_pptx"
    description = "Read a .pptx file's slide text, speaker notes, and table text."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "slide_range": {
                "type": "array",
                "items": {"type": "integer"},
                "minItems": 2,
                "maxItems": 2,
                "description": "[start, end] slide indices, 0-based inclusive.",
            },
        },
        "required": ["path"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")))
        slide_range = args.get("slide_range")
        return await asyncio.to_thread(self._read, resolved, slide_range)

    def _read(self, path: Path, slide_range: object) -> ToolResult:
        if not path.exists():
            return ToolResult.text(f"Not found: {path}", is_error=True)
        try:
            presentation = pptx.Presentation(str(path))
        except Exception as exc:  # the file may be corrupt, encrypted, or not a pptx
            return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)

        slides = list(presentation.slides)
        start, end = 0, len(slides) - 1
        if isinstance(slide_range, list) and len(slide_range) == 2:
            start, end = int(slide_range[0]), int(slide_range[1])

        lines: list[str] = []
        for index, slide in enumerate(slides):
            if index < start or index > end:
                continue
            lines.append(f"\n--- Slide {index} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text.strip() for cell in row.cells))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    lines.append(f"[notes] {notes}")

        return ToolResult.text(truncate("\n".join(lines) if lines else "(no content)"))


def _iter_all_text_frames(presentation: Any) -> Iterator[Any]:
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield shape.text_frame
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame


def _replace_text_preserving_runs(presentation: Any, old: str, new: str) -> None:
    if not old:
        return
    for text_frame in _iter_all_text_frames(presentation):
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def _add_text_slide(presentation: Any, title: str, bullets: list[str]) -> None:
    layouts = presentation.slide_layouts
    layout = layouts[1] if len(layouts) > 1 else layouts[0]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1 and bullets:
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet


def _as_str_list(value: object) -> list[str]:
    return [str(v) for v in value] if isinstance(value, list) else []


def _preview_ops(presentation: Any, ops: list[dict[str, Any]]) -> list[str]:
    lines = []
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            old, new = str(op.get("old", "")), str(op.get("new", ""))
            count = (
                sum(tf.text.count(old) for tf in _iter_all_text_frames(presentation)) if old else 0
            )
            lines.append(f"replace_text: {old!r} → {new!r} ({count} occurrence(s))")
        elif kind == "set_notes":
            lines.append(f"set_notes: slide {op.get('slide_index')} notes updated")
        elif kind == "add_text_slide":
            bullets = _as_str_list(op.get("bullets"))
            lines.append(f"add_text_slide: + {op.get('title', '')!r} ({len(bullets)} bullet(s))")
        else:
            lines.append(f"unknown op: {kind!r}")
    return lines


def _apply_ops(presentation: Any, ops: list[dict[str, Any]]) -> None:
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            _replace_text_preserving_runs(
                presentation, str(op.get("old", "")), str(op.get("new", ""))
            )
        elif kind == "set_notes":
            slide_index = as_int(op.get("slide_index"), 0)
            presentation.slides[slide_index].notes_slide.notes_text_frame.text = str(
                op.get("text", "")
            )
        elif kind == "add_text_slide":
            _add_text_slide(presentation, str(op.get("title", "")), _as_str_list(op.get("bullets")))


class UpdatePptxTool(Tool):
    name = "update_pptx"
    description = (
        "Apply ops to a .pptx file (created fresh if it doesn't exist yet): replace_text, "
        "set_notes, add_text_slide."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "ops": {"type": "array", "items": {"type": "object"}},
        },
        "required": ["path", "ops"],
    }
    risk = RiskLevel.WRITE

    async def preview(self, args: dict[str, object], ctx: ToolContext) -> str:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        ops = _as_ops_list(args.get("ops"))
        return await asyncio.to_thread(self._preview, resolved, ops)

    def _preview(self, path: Path, ops: list[dict[str, Any]]) -> str:
        try:
            presentation = pptx.Presentation(str(path)) if path.exists() else pptx.Presentation()
        except Exception as exc:  # the file may be corrupt, encrypted, or not a pptx
            return f"Could not open {path}: {exc}"
        header = f"Create new file: {path}\n" if not path.exists() else ""
        return header + ("\n".join(_preview_ops(presentation, ops)) or "(no ops)")

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        ops = _as_ops_list(args.get("ops"))
        return await asyncio.to_thread(self._update, resolved, ops)

    def _update(self, path: Path, ops: list[dict[str, Any]]) -> ToolResult:
        try:
            presentation = pptx.Presentation(str(path)) if path.exists() else pptx.Presentation()
        except Exception as exc:  # the file may be corrupt, encrypted, or not a pptx
            return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)
        backup_existing(path)
        try:
            _apply_ops(presentation, ops)
        except Exception as exc:  # a malformed op must not crash the agent loop
            return ToolResult.text(f"Failed applying ops: {exc}", is_error=True)
        path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(path))
        return ToolResult.text(f"Applied {len(ops)} op(s) to {path}")


def _as_ops_list(value: object) -> list[dict[str, Any]]:
    return [op for op in value if isinstance(op, dict)] if isinstance(value, list) else []
