"""Excel workbook tools (ARCHITECTURE.md §8.2): read, and write/update cells."""

from __future__ import annotations

import asyncio
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from deskai.tools.base import RiskLevel, Tool, ToolContext, ToolResult, backup_existing, truncate

_MAX_ROWS = 200
_MAX_COLS = 50
_ISO_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")


def _as_cell_value(value: Any) -> Any:
    """A JSON tool call can't send a real date — a plain `YYYY-MM-DD` string is
    the closest it gets, so that shape is turned into an actual date here,
    which openpyxl then renders as a real Excel date rather than inert text.

    openpyxl always reads a date-formatted cell back as `datetime` (xlsx has
    no separate date type at the file-format level), so this returns a
    `datetime` too — a plain `date` would never compare equal to what a
    saved-and-reread cell holds, breaking the write/preview no-op check.
    """
    if isinstance(value, str) and _ISO_DATE_RE.fullmatch(value):
        try:
            return datetime.combine(date.fromisoformat(value), datetime.min.time())
        except ValueError:
            return value
    return value


def _rows_for_range(worksheet: Any, cell_range: str) -> Any:
    """openpyxl ships no type stubs, so this stays loosely typed at the boundary."""
    selected = worksheet[cell_range]
    if isinstance(selected, Cell):
        return ((selected,),)
    if selected and isinstance(selected[0], Cell):
        return (selected,)
    return selected


class ReadXlsxTool(Tool):
    name = "read_xlsx"
    description = (
        "Read values from an .xlsx sheet (a range, or the whole sheet), or list sheet names."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {
                "type": "string",
                "description": "Sheet name; omit to list available sheets.",
            },
            "range": {"type": "string", "description": "A1-style range, e.g. 'A1:D20'."},
        },
        "required": ["path"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")))
        sheet = args.get("sheet")
        cell_range = args.get("range")
        return await asyncio.to_thread(
            self._read,
            resolved,
            str(sheet) if sheet else None,
            str(cell_range) if cell_range else None,
        )

    def _read(self, path: Path, sheet_name: str | None, cell_range: str | None) -> ToolResult:
        if not path.exists():
            return ToolResult.text(f"Not found: {path}", is_error=True)
        try:
            workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        except Exception as exc:  # the file may be corrupt, encrypted, or not an xlsx
            return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)

        try:
            if sheet_name is None:
                return ToolResult.text("Sheets: " + ", ".join(workbook.sheetnames))
            if sheet_name not in workbook.sheetnames:
                available = ", ".join(workbook.sheetnames)
                return ToolResult.text(
                    f"No such sheet: {sheet_name!r}. Sheets: {available}", is_error=True
                )

            worksheet = workbook[sheet_name]
            if cell_range:
                rows = _rows_for_range(worksheet, cell_range)
            else:
                rows = tuple(worksheet.iter_rows(max_row=_MAX_ROWS, max_col=_MAX_COLS))

            lines = [
                " | ".join("" if cell.value is None else str(cell.value) for cell in row)
                for row in rows
            ]
            body = "\n".join(lines) if lines else "(empty)"
            return ToolResult.text(truncate(f"{sheet_name}\n{body}"))
        finally:
            workbook.close()


def _anchor_from_range(cell_range: str) -> tuple[int, int]:
    col_letter, row = coordinate_from_string(cell_range.split(":")[0])
    return column_index_from_string(col_letter), row


class WriteXlsxTool(Tool):
    name = "write_xlsx"
    description = (
        "Create or update cells in an .xlsx sheet, starting at a range's top-left cell. "
        "Keeps all other content in the file."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string"},
            "range": {"type": "string", "description": "Top-left anchor cell, e.g. 'A1'."},
            "values": {
                "type": "array",
                "items": {"type": "array"},
                "description": "Rows of cell values.",
            },
        },
        "required": ["path", "sheet", "range", "values"],
    }
    risk = RiskLevel.WRITE

    def _parse(self, args: dict[str, object]) -> tuple[str, int, int, list[list[Any]]]:
        sheet_name = str(args.get("sheet") or "Sheet1")
        anchor_col, anchor_row = _anchor_from_range(str(args.get("range") or "A1"))
        raw_values = args.get("values")
        values = [list(row) for row in raw_values] if isinstance(raw_values, list) else []
        return sheet_name, anchor_col, anchor_row, values

    async def preview(self, args: dict[str, object], ctx: ToolContext) -> str:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        sheet_name, anchor_col, anchor_row, values = self._parse(args)
        return await asyncio.to_thread(
            self._preview, resolved, sheet_name, anchor_col, anchor_row, values
        )

    def _preview(
        self,
        path: Path,
        sheet_name: str,
        anchor_col: int,
        anchor_row: int,
        values: list[list[Any]],
    ) -> str:
        cell_count = sum(len(row) for row in values)
        if not path.exists():
            return f"Create new file {path} — sheet {sheet_name!r} with {cell_count} cell(s)"

        try:
            workbook = openpyxl.load_workbook(str(path), data_only=True)
        except Exception as exc:  # the file may be corrupt, encrypted, or not an xlsx
            return f"Could not open {path}: {exc}"
        try:
            if sheet_name not in workbook.sheetnames:
                return f"Create new sheet {sheet_name!r} with {cell_count} cell(s) in {path}"
            worksheet = workbook[sheet_name]
            changes = []
            for r, row in enumerate(values):
                for c, value in enumerate(row):
                    cell = worksheet.cell(row=anchor_row + r, column=anchor_col + c)
                    new_value = _as_cell_value(value)
                    if cell.value != new_value:
                        changes.append(
                            f"{sheet_name}!{cell.coordinate}: {cell.value!r} → {new_value!r}"
                        )
            return "\n".join(changes) if changes else "(no changes)"
        finally:
            workbook.close()

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        sheet_name, anchor_col, anchor_row, values = self._parse(args)
        return await asyncio.to_thread(
            self._write, resolved, sheet_name, anchor_col, anchor_row, values
        )

    def _write(
        self,
        path: Path,
        sheet_name: str,
        anchor_col: int,
        anchor_row: int,
        values: list[list[Any]],
    ) -> ToolResult:
        backup_existing(path)
        if path.exists():
            try:
                workbook = openpyxl.load_workbook(str(path))
            except Exception as exc:  # the file may be corrupt, encrypted, or not an xlsx
                return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)
        else:
            workbook = openpyxl.Workbook()
            default_sheet = workbook.active
            if default_sheet is not None:
                default_sheet.title = sheet_name

        if sheet_name not in workbook.sheetnames:
            workbook.create_sheet(sheet_name)
        worksheet = workbook[sheet_name]

        for r, row in enumerate(values):
            for c, value in enumerate(row):
                cell = worksheet.cell(row=anchor_row + r, column=anchor_col + c)
                cell.value = _as_cell_value(value)
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(path))
        workbook.close()
        cell_count = sum(len(row) for row in values)
        return ToolResult.text(f"Wrote {cell_count} cell(s) to {sheet_name!r} in {path}")
