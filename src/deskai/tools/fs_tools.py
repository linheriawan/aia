"""Filesystem read tools (ARCHITECTURE.md §8.2): list_dir, file_info, search_files."""

from __future__ import annotations

import asyncio
import fnmatch
from datetime import datetime
from pathlib import Path

from deskai.tools.base import RiskLevel, Tool, ToolContext, ToolResult, as_int, truncate

_MAX_ENTRIES = 500
_MAX_SEARCH_MATCHES = 200
_TEXT_SUFFIXES = {".txt", ".md", ".csv", ".log"}


class ListDirTool(Tool):
    name = "list_dir"
    description = "List files and subfolders under a workspace path, with size and modified time."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Relative or absolute path to list."},
            "depth": {"type": "integer", "minimum": 1, "maximum": 5},
            "pattern": {"type": "string", "description": "Optional filename glob, e.g. '*.pdf'."},
        },
        "required": ["path"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        root = ctx.workspace.resolve(str(args.get("path", ".")))
        depth = as_int(args.get("depth"), default=1)
        pattern = args.get("pattern")
        return await asyncio.to_thread(self._list, root, depth, str(pattern) if pattern else None)

    def _list(self, root: Path, depth: int, pattern: str | None) -> ToolResult:
        if not root.exists():
            return ToolResult.text(f"Not found: {root}", is_error=True)
        if not root.is_dir():
            return ToolResult.text(f"Not a directory: {root}", is_error=True)

        lines: list[str] = []
        count = 0

        def describe(entry: Path) -> str:
            try:
                stat = entry.stat()
                modified = datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds")
                return f"{stat.st_size}B, {modified}"
            except OSError:
                return "?"

        def walk(directory: Path, remaining_depth: int, prefix: str) -> None:
            nonlocal count
            try:
                entries = sorted(directory.iterdir(), key=lambda p: (p.is_file(), p.name.lower()))
            except OSError as exc:
                lines.append(f"{prefix}[error reading {directory.name}: {exc}]")
                return
            for entry in entries:
                if count >= _MAX_ENTRIES:
                    lines.append(f"{prefix}… ({_MAX_ENTRIES}+ entries, capped)")
                    return
                if entry.is_dir():
                    lines.append(f"{prefix}{entry.name}/")
                    count += 1
                    if remaining_depth > 1:
                        walk(entry, remaining_depth - 1, prefix + "  ")
                else:
                    if pattern and not fnmatch.fnmatch(entry.name, pattern):
                        continue
                    lines.append(f"{prefix}{entry.name}  [{describe(entry)}]")
                    count += 1

        walk(root, depth, "")
        return ToolResult.text(truncate("\n".join(lines) if lines else "(empty)"))


class FileInfoTool(Tool):
    name = "file_info"
    description = (
        "Get a file's type, size, and (for recognized document types) page/sheet/slide count."
    )
    input_schema = {
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")))
        return await asyncio.to_thread(self._info, resolved)

    def _info(self, path: Path) -> ToolResult:
        if not path.exists():
            return ToolResult.text(f"Not found: {path}", is_error=True)

        suffix = path.suffix.lower()
        lines = [
            f"Path: {path}",
            f"Type: {'directory' if path.is_dir() else (suffix or 'file')}",
            f"Size: {path.stat().st_size} bytes",
        ]
        if path.is_dir():
            return ToolResult.text("\n".join(lines))

        try:
            if suffix == ".docx":
                import docx

                document = docx.Document(str(path))
                lines.append(f"Paragraphs: {len(document.paragraphs)}")
                lines.append(f"Tables: {len(document.tables)}")
            elif suffix == ".xlsx":
                import openpyxl

                workbook = openpyxl.load_workbook(str(path), read_only=True)
                lines.append(f"Sheets: {', '.join(workbook.sheetnames)}")
                workbook.close()
            elif suffix == ".pptx":
                import pptx

                presentation = pptx.Presentation(str(path))
                lines.append(f"Slides: {len(presentation.slides)}")
            elif suffix == ".pdf":
                import pymupdf

                with pymupdf.open(str(path)) as document:
                    lines.append(f"Pages: {document.page_count}")
        except Exception as exc:  # the file may be corrupt, encrypted, or unrecognized
            lines.append(f"(could not read document metadata: {exc})")
        return ToolResult.text("\n".join(lines))


class SearchFilesTool(Tool):
    name = "search_files"
    description = "Search filenames, and the content of plain-text files, under a workspace path."
    input_schema = {
        "type": "object",
        "properties": {
            "query": {"type": "string"},
            "path": {
                "type": "string",
                "description": "Folder to search; defaults to the workspace root.",
            },
            "glob": {"type": "string", "description": "Optional filename glob, e.g. '*.txt'."},
        },
        "required": ["query"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        query = str(args.get("query", "")).strip()
        if not query:
            return ToolResult.text("A search query is required.", is_error=True)
        root = ctx.workspace.resolve(str(args.get("path") or "."))
        glob = args.get("glob")
        return await asyncio.to_thread(self._search, root, query, str(glob) if glob else None)

    def _search(self, root: Path, query: str, glob: str | None) -> ToolResult:
        if not root.exists():
            return ToolResult.text(f"Not found: {root}", is_error=True)

        query_lower = query.lower()
        matches: list[str] = []
        for file_path in sorted(root.rglob(glob or "*")):
            if len(matches) >= _MAX_SEARCH_MATCHES:
                matches.append(f"… ({_MAX_SEARCH_MATCHES}+ matches, capped)")
                break
            if not file_path.is_file():
                continue

            name_match = query_lower in file_path.name.lower()
            content_line = None
            if file_path.suffix.lower() in _TEXT_SUFFIXES:
                content_line = self._first_matching_line(file_path, query_lower)

            if name_match or content_line:
                relative = file_path.relative_to(root)
                entry = str(relative) + (f"  ({content_line})" if content_line else "")
                matches.append(entry)

        return ToolResult.text(truncate("\n".join(matches) if matches else "No matches."))

    @staticmethod
    def _first_matching_line(file_path: Path, query_lower: str) -> str | None:
        try:
            with file_path.open("r", encoding="utf-8", errors="replace") as handle:
                for line_number, line in enumerate(handle, start=1):
                    if query_lower in line.lower():
                        return f"line {line_number}: {line.strip()[:200]}"
        except OSError:
            return None
        return None
