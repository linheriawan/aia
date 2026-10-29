"""Word document tools (ARCHITECTURE.md §8.2): read, create from blocks, and update by ops."""

from __future__ import annotations

import asyncio
from collections import Counter
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import docx

from deskai.tools.base import (
    RiskLevel,
    Tool,
    ToolContext,
    ToolResult,
    as_int,
    backup_existing,
    truncate,
)


def _table_to_markdown(table: Any) -> list[str]:
    rows = table.rows
    if not rows:
        return []
    lines = ["| " + " | ".join(c.text.strip() for c in rows[0].cells) + " |"]
    lines.append("| " + " | ".join("---" for _ in rows[0].cells) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    return lines


class ReadDocxTool(Tool):
    name = "read_docx"
    description = "Read a .docx file's paragraphs (with style) and tables as markdown."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "paragraph_range": {
                "type": "array",
                "items": {"type": "integer"},
                "minItems": 2,
                "maxItems": 2,
                "description": "[start, end] paragraph indices, 0-based inclusive.",
            },
        },
        "required": ["path"],
    }
    risk = RiskLevel.READ

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")))
        paragraph_range = args.get("paragraph_range")
        return await asyncio.to_thread(self._read, resolved, paragraph_range)

    def _read(self, path: Path, paragraph_range: object) -> ToolResult:
        if not path.exists():
            return ToolResult.text(f"Not found: {path}", is_error=True)
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # the file may be corrupt, encrypted, or not a docx
            return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)

        start, end = 0, len(document.paragraphs) - 1
        if isinstance(paragraph_range, list) and len(paragraph_range) == 2:
            start, end = int(paragraph_range[0]), int(paragraph_range[1])

        lines: list[str] = []
        for index, paragraph in enumerate(document.paragraphs):
            if start <= index <= end and paragraph.text.strip():
                style = paragraph.style.name if paragraph.style else "Normal"
                lines.append(f"[{index}] ({style}) {paragraph.text}")

        for table_index, table in enumerate(document.tables):
            table_lines = _table_to_markdown(table)
            if table_lines:
                lines.append(f"\nTable {table_index}:")
                lines.extend(table_lines)

        return ToolResult.text(truncate("\n".join(lines) if lines else "(no content)"))


def _apply_blocks(document: Any, blocks: list[dict[str, Any]]) -> None:
    for block in blocks:
        kind = block.get("type")
        if kind == "heading":
            document.add_heading(str(block.get("text", "")), level=as_int(block.get("level"), 1))
        elif kind == "paragraph":
            document.add_paragraph(str(block.get("text", "")))
        elif kind == "bullet":
            document.add_paragraph(str(block.get("text", "")), style="List Bullet")
        elif kind == "table":
            rows = block.get("rows") or []
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.rows[r].cells[c].text = str(value)


class CreateDocxTool(Tool):
    name = "create_docx"
    description = (
        "Create a new Word document from blocks (heading/paragraph/bullet/table). "
        "Refuses if the file already exists — use update_docx for an existing one."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "blocks": {
                "type": "array",
                "items": {"type": "object"},
                "description": "Each block: {type: heading|paragraph|bullet|table, ...}",
            },
        },
        "required": ["path", "blocks"],
    }
    risk = RiskLevel.WRITE

    async def preview(self, args: dict[str, object], ctx: ToolContext) -> str:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        blocks = _as_block_list(args.get("blocks"))
        if resolved.exists():
            return f"{resolved} already exists — create_docx refuses to overwrite it."
        counts = Counter(str(b.get("type")) for b in blocks)
        summary = ", ".join(f"{n} {t}" for t, n in counts.items())
        return f"Create new Word document: {resolved}\n{summary or 'no blocks'}"

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        blocks = _as_block_list(args.get("blocks"))
        return await asyncio.to_thread(self._create, resolved, blocks)

    def _create(self, path: Path, blocks: list[dict[str, Any]]) -> ToolResult:
        if path.exists():
            return ToolResult.text(
                f"{path} already exists; use update_docx instead.", is_error=True
            )
        document = docx.Document()
        _apply_blocks(document, blocks)
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(path))
        return ToolResult.text(f"Created {path} with {len(blocks)} block(s)")


def _as_block_list(value: object) -> list[dict[str, Any]]:
    return [b for b in value if isinstance(b, dict)] if isinstance(value, list) else []


def _iter_all_paragraphs(document: Any) -> Iterator[Any]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_text_preserving_runs(document: Any, old: str, new: str) -> None:
    if not old:
        return
    for paragraph in _iter_all_paragraphs(document):
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def _insert_paragraph_after(document: Any, index: int, text: str) -> None:
    paragraphs = document.paragraphs
    if 0 <= index + 1 < len(paragraphs):
        paragraphs[index + 1].insert_paragraph_before(text)
    else:
        document.add_paragraph(text)


def _preview_ops(document: Any, ops: list[dict[str, Any]]) -> list[str]:
    lines = []
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            old, new = str(op.get("old", "")), str(op.get("new", ""))
            count = sum(p.text.count(old) for p in _iter_all_paragraphs(document)) if old else 0
            lines.append(f"replace_text: {old!r} → {new!r} ({count} occurrence(s))")
        elif kind == "append_paragraph":
            lines.append(f"append_paragraph: + {op.get('text', '')!r}")
        elif kind == "insert_paragraph_after":
            lines.append(f"insert_paragraph_after[{op.get('index')}]: + {op.get('text', '')!r}")
        elif kind == "set_table_cell":
            table_index = as_int(op.get("table_index"), 0)
            row, col = as_int(op.get("row"), 0), as_int(op.get("col"), 0)
            try:
                old_value = document.tables[table_index].cell(row, col).text
            except (IndexError, KeyError):
                old_value = "?"
            new_value = op.get("text", "")
            lines.append(
                f"set_table_cell: Table{table_index}[{row}][{col}]: {old_value!r} → {new_value!r}"
            )
        else:
            lines.append(f"unknown op: {kind!r}")
    return lines


def _apply_ops(document: Any, ops: list[dict[str, Any]]) -> None:
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            _replace_text_preserving_runs(document, str(op.get("old", "")), str(op.get("new", "")))
        elif kind == "append_paragraph":
            document.add_paragraph(str(op.get("text", "")))
        elif kind == "insert_paragraph_after":
            _insert_paragraph_after(document, as_int(op.get("index"), 0), str(op.get("text", "")))
        elif kind == "set_table_cell":
            table_index = as_int(op.get("table_index"), 0)
            row, col = as_int(op.get("row"), 0), as_int(op.get("col"), 0)
            document.tables[table_index].cell(row, col).text = str(op.get("text", ""))


class UpdateDocxTool(Tool):
    name = "update_docx"
    description = (
        "Apply ops to an existing Word document: replace_text, append_paragraph, "
        "insert_paragraph_after, set_table_cell. Preserves run formatting where possible."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "ops": {"type": "array", "items": {"type": "object"}},
        },
        "required": ["path", "ops"],
    }
    risk = RiskLevel.WRITE

    async def preview(self, args: dict[str, object], ctx: ToolContext) -> str:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        ops = _as_block_list(args.get("ops"))
        return await asyncio.to_thread(self._preview, resolved, ops)

    def _preview(self, path: Path, ops: list[dict[str, Any]]) -> str:
        if not path.exists():
            return f"Not found: {path}"
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # the file may be corrupt, encrypted, or not a docx
            return f"Could not open {path}: {exc}"
        return "\n".join(_preview_ops(document, ops)) or "(no ops)"

    async def run(self, args: dict[str, object], ctx: ToolContext) -> ToolResult:
        resolved = ctx.workspace.resolve(str(args.get("path", "")), for_write=True)
        ops = _as_block_list(args.get("ops"))
        return await asyncio.to_thread(self._update, resolved, ops)

    def _update(self, path: Path, ops: list[dict[str, Any]]) -> ToolResult:
        if not path.exists():
            return ToolResult.text(f"Not found: {path}", is_error=True)
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # the file may be corrupt, encrypted, or not a docx
            return ToolResult.text(f"Could not open {path}: {exc}", is_error=True)
        backup_existing(path)
        try:
            _apply_ops(document, ops)
        except Exception as exc:  # a malformed op must not crash the agent loop
            return ToolResult.text(f"Failed applying ops: {exc}", is_error=True)
        document.save(str(path))
        return ToolResult.text(f"Applied {len(ops)} op(s) to {path}")
