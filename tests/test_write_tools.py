"""Tests for the Phase 4 write tools: creation, update-preserving-other-content,
backups, and the read-only-folder rejection (CLAUDE.md hard rules)."""

from __future__ import annotations

import asyncio
from datetime import date, datetime
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pptx
import pytest

from deskai.ocr.engine import OcrEngine
from deskai.storage import db
from deskai.tools.base import ToolContext, ToolResult
from deskai.tools.docx_tools import CreateDocxTool, UpdateDocxTool
from deskai.tools.pptx_tools import UpdatePptxTool
from deskai.tools.sandbox import Sandbox, SandboxError, WorkspaceFolder
from deskai.tools.text_tools import EditTextTool, WriteTextTool
from deskai.tools.xlsx_tools import WriteXlsxTool


def _run(coro: Any) -> Any:
    return asyncio.run(coro)


@pytest.fixture
def workspace(tmp_path: Path) -> Path:
    return tmp_path


@pytest.fixture
def ctx(workspace: Path, tmp_path: Path) -> ToolContext:
    sandbox = Sandbox([WorkspaceFolder(workspace)], deny_list=[])
    conn = db.connect(tmp_path / "test.db")
    return ToolContext(session_id="s1", workspace=sandbox, ocr=OcrEngine(conn))


@pytest.fixture
def readonly_ctx(workspace: Path, tmp_path: Path) -> ToolContext:
    sandbox = Sandbox([WorkspaceFolder(workspace, read_only=True)], deny_list=[])
    conn = db.connect(tmp_path / "test2.db")
    return ToolContext(session_id="s1", workspace=sandbox, ocr=OcrEngine(conn))


# --- write_text / edit_text --------------------------------------------------


def test_write_text_creates_new_file(ctx: ToolContext, workspace: Path) -> None:
    result: ToolResult = _run(WriteTextTool().run({"path": "note.txt", "content": "hello"}, ctx))
    assert not result.is_error
    assert (workspace / "note.txt").read_text() == "hello"


def test_write_text_refuses_overwrite_without_flag(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "note.txt").write_text("original")
    result = _run(WriteTextTool().run({"path": "note.txt", "content": "new"}, ctx))
    assert result.is_error
    assert (workspace / "note.txt").read_text() == "original"


def test_write_text_overwrite_creates_backup(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "note.txt").write_text("original")
    result = _run(
        WriteTextTool().run({"path": "note.txt", "content": "new", "overwrite": True}, ctx)
    )
    assert not result.is_error
    assert (workspace / "note.txt").read_text() == "new"
    backups = list(workspace.glob("note.txt.bak-*"))
    assert len(backups) == 1
    assert backups[0].read_text() == "original"


def test_write_text_preview_shows_diff(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "note.txt").write_text("old line\n")
    preview = _run(
        WriteTextTool().preview(
            {"path": "note.txt", "content": "new line\n", "overwrite": True}, ctx
        )
    )
    assert "-old line" in preview
    assert "+new line" in preview


def test_write_text_rejected_in_read_only_folder(readonly_ctx: ToolContext) -> None:
    with pytest.raises(SandboxError):
        _run(WriteTextTool().run({"path": "note.txt", "content": "hi"}, readonly_ctx))


def test_edit_text_requires_exactly_one_match(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "note.txt").write_text("foo foo")
    result = _run(EditTextTool().run({"path": "note.txt", "old": "foo", "new": "bar"}, ctx))
    assert result.is_error
    assert (workspace / "note.txt").read_text() == "foo foo"


def test_edit_text_replaces_single_match_with_backup(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "note.txt").write_text("hello world")
    result = _run(EditTextTool().run({"path": "note.txt", "old": "world", "new": "there"}, ctx))
    assert not result.is_error
    assert (workspace / "note.txt").read_text() == "hello there"
    assert list(workspace.glob("note.txt.bak-*"))


# --- create_docx / update_docx -----------------------------------------------


def test_create_docx_builds_blocks(ctx: ToolContext, workspace: Path) -> None:
    blocks = [
        {"type": "heading", "text": "Title", "level": 1},
        {"type": "paragraph", "text": "Body text"},
        {"type": "bullet", "text": "Bullet one"},
        {"type": "table", "rows": [["A", "B"], ["1", "2"]]},
    ]
    result = _run(CreateDocxTool().run({"path": "report.docx", "blocks": blocks}, ctx))
    assert not result.is_error

    document = docx.Document(str(workspace / "report.docx"))
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "Title" in all_text
    assert "Body text" in all_text
    assert "Bullet one" in all_text
    assert document.tables[0].cell(1, 1).text == "2"


def test_create_docx_refuses_to_overwrite(ctx: ToolContext, workspace: Path) -> None:
    (workspace / "report.docx").write_bytes(b"not a real docx, but exists")
    result = _run(CreateDocxTool().run({"path": "report.docx", "blocks": []}, ctx))
    assert result.is_error


def test_update_docx_replace_text_preserves_other_runs(ctx: ToolContext, workspace: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph()
    run1 = paragraph.add_run("Hello ")
    run1.bold = True
    paragraph.add_run("World")
    document.save(str(workspace / "doc.docx"))

    ops = [{"op": "replace_text", "old": "World", "new": "Earth"}]
    result = _run(UpdateDocxTool().run({"path": "doc.docx", "ops": ops}, ctx))
    assert not result.is_error

    updated = docx.Document(str(workspace / "doc.docx"))
    runs = updated.paragraphs[0].runs
    assert runs[0].text == "Hello "
    assert runs[0].bold is True
    assert runs[1].text == "Earth"
    assert list(workspace.glob("doc.docx.bak-*"))


def test_update_docx_append_and_insert_paragraph(ctx: ToolContext, workspace: Path) -> None:
    document = docx.Document()
    document.add_paragraph("First")
    document.add_paragraph("Last")
    document.save(str(workspace / "doc.docx"))

    ops = [
        {"op": "insert_paragraph_after", "index": 0, "text": "Middle"},
        {"op": "append_paragraph", "text": "Appended"},
    ]
    result = _run(UpdateDocxTool().run({"path": "doc.docx", "ops": ops}, ctx))
    assert not result.is_error

    updated = docx.Document(str(workspace / "doc.docx"))
    texts = [p.text for p in updated.paragraphs]
    assert texts == ["First", "Middle", "Last", "Appended"]


def test_update_docx_set_table_cell(ctx: ToolContext, workspace: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "old"
    document.save(str(workspace / "doc.docx"))

    ops = [{"op": "set_table_cell", "table_index": 0, "row": 0, "col": 0, "text": "new"}]
    result = _run(UpdateDocxTool().run({"path": "doc.docx", "ops": ops}, ctx))
    assert not result.is_error

    updated = docx.Document(str(workspace / "doc.docx"))
    assert updated.tables[0].cell(0, 0).text == "new"


def test_update_docx_preview_describes_changes_without_writing(
    ctx: ToolContext, workspace: Path
) -> None:
    document = docx.Document()
    document.add_paragraph("Hello World")
    document.save(str(workspace / "doc.docx"))
    original_bytes = (workspace / "doc.docx").read_bytes()

    ops = [{"op": "replace_text", "old": "World", "new": "Earth"}]
    preview = _run(UpdateDocxTool().preview({"path": "doc.docx", "ops": ops}, ctx))

    assert "World" in preview
    assert "Earth" in preview
    assert (workspace / "doc.docx").read_bytes() == original_bytes  # untouched


# --- write_xlsx ---------------------------------------------------------------


def test_write_xlsx_creates_new_file_and_sheet(ctx: ToolContext, workspace: Path) -> None:
    result = _run(
        WriteXlsxTool().run(
            {
                "path": "data.xlsx",
                "sheet": "Data",
                "range": "A1",
                "values": [["Name", "Amount"], ["Widget", 100]],
            },
            ctx,
        )
    )
    assert not result.is_error

    workbook = openpyxl.load_workbook(str(workspace / "data.xlsx"))
    sheet = workbook["Data"]
    assert sheet["A1"].value == "Name"
    assert sheet["B2"].value == 100


def test_write_xlsx_keeps_other_cells(ctx: ToolContext, workspace: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Keep me"
    sheet["C1"] = "Also keep"
    workbook.save(str(workspace / "data.xlsx"))

    result = _run(
        WriteXlsxTool().run(
            {"path": "data.xlsx", "sheet": "Data", "range": "B1", "values": [["New"]]}, ctx
        )
    )
    assert not result.is_error

    reopened = openpyxl.load_workbook(str(workspace / "data.xlsx"))
    sheet = reopened["Data"]
    assert sheet["A1"].value == "Keep me"
    assert sheet["B1"].value == "New"
    assert sheet["C1"].value == "Also keep"


def test_write_xlsx_preview_matches_architecture_example_format(
    ctx: ToolContext, workspace: Path
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["B2"] = 100
    workbook.save(str(workspace / "data.xlsx"))

    preview = _run(
        WriteXlsxTool().preview(
            {"path": "data.xlsx", "sheet": "Sheet1", "range": "B2", "values": [[120]]}, ctx
        )
    )
    assert "Sheet1!B2: 100 → 120" in preview


def test_write_xlsx_converts_iso_date_strings_to_real_dates(
    ctx: ToolContext, workspace: Path
) -> None:
    result = _run(
        WriteXlsxTool().run(
            {
                "path": "data.xlsx",
                "sheet": "Data",
                "range": "A1",
                "values": [["2026-10-23", "not a date"]],
            },
            ctx,
        )
    )
    assert not result.is_error

    workbook = openpyxl.load_workbook(str(workspace / "data.xlsx"))
    sheet = workbook["Data"]
    assert sheet["A1"].value == datetime(2026, 10, 23)
    assert sheet["A1"].number_format == "yyyy-mm-dd"
    assert sheet["B1"].value == "not a date"


def test_write_xlsx_preview_treats_iso_date_strings_as_dates(
    ctx: ToolContext, workspace: Path
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = date(2026, 10, 23)
    workbook.save(str(workspace / "data.xlsx"))

    preview = _run(
        WriteXlsxTool().preview(
            {"path": "data.xlsx", "sheet": "Sheet1", "range": "A1", "values": [["2026-10-23"]]},
            ctx,
        )
    )
    assert preview == "(no changes)"


def test_write_xlsx_rejected_in_read_only_folder(readonly_ctx: ToolContext) -> None:
    with pytest.raises(SandboxError):
        _run(
            WriteXlsxTool().run(
                {"path": "data.xlsx", "sheet": "S", "range": "A1", "values": [[1]]},
                readonly_ctx,
            )
        )


# --- update_pptx ---------------------------------------------------------------


def test_update_pptx_creates_file_when_missing(ctx: ToolContext, workspace: Path) -> None:
    ops = [{"op": "add_text_slide", "title": "New Slide", "bullets": ["one", "two"]}]
    result = _run(UpdatePptxTool().run({"path": "deck.pptx", "ops": ops}, ctx))
    assert not result.is_error

    presentation = pptx.Presentation(str(workspace / "deck.pptx"))
    assert len(presentation.slides) == 1
    assert presentation.slides[0].shapes.title.text == "New Slide"


def test_update_pptx_replace_text_and_set_notes(ctx: ToolContext, workspace: Path) -> None:
    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Old Title"
    presentation.save(str(workspace / "deck.pptx"))

    ops = [
        {"op": "replace_text", "old": "Old Title", "new": "New Title"},
        {"op": "set_notes", "slide_index": 0, "text": "Speaker notes here"},
    ]
    result = _run(UpdatePptxTool().run({"path": "deck.pptx", "ops": ops}, ctx))
    assert not result.is_error

    updated = pptx.Presentation(str(workspace / "deck.pptx"))
    assert updated.slides[0].shapes.title.text == "New Title"
    assert updated.slides[0].notes_slide.notes_text_frame.text == "Speaker notes here"
    assert list(workspace.glob("deck.pptx.bak-*"))
